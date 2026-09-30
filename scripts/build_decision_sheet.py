from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
ROOT=Path(__file__).resolve().parents[1]
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.5);sec.left_margin=sec.right_margin=Inches(.65)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11);style.paragraph_format.space_after=Pt(5)
for name in ['Heading 1','Heading 2']:
 doc.styles[name].font.color.rgb=RGBColor.from_string('083661');doc.styles[name].paragraph_format.space_after=Pt(6)
doc.add_heading('From a forecast to an order',0)
doc.add_paragraph('Sina Mirshahi · JKU Linz · 03.10.2026\nWork in pairs. Keep this as your own notes.')
doc.add_heading('1. Choose a forecast',2)
for t in ['Product and channel: __________________________________________',
          'Window:  Spring / Summer / Christmas     Method: ___________________',
          'Evidence: average miss ______ units; bias ______ units.',
          'A different season or channel changed our choice because: __________',
          'Before trusting this result, we would check: ________________________']:
 doc.add_paragraph(t)
doc.add_heading('2. Calculate an order',2)
for t in ['Five weekly forecasts: 180 + 190 + 210 + 200 + 220 = ______ units.',
          'Add safety stock of 148: stock target = ______ units.',
          'Stock on shelf = 300; already on order = 500.',
          'Order now = stock target minus those 800 units = ______ units.']:
 doc.add_paragraph(t)
p=doc.add_paragraph('Round teaching numbers. Assume no backorders, fixed lead time and scheduled arrivals.');p.runs[0].italic=True;p.runs[0].font.size=Pt(9)
doc.add_heading('3. Choose a stock policy',2)
doc.add_paragraph('Product: ____________________________________________________')
t=doc.add_table(rows=3,cols=5);t.style='Light Shading Accent 1'
for c,text in zip(t.rows[0].cells,['Setting','Cycle service target','Average stock, units','Fill rate','Stockout weeks']):c.text=text
for row,label in zip(t.rows[1:],['A','B']):row.cells[0].text=label;row.height=Inches(.36)
for text in ['We would trial setting ______ because: ___________________________',
             'We accept this consequence: ___________________________________',
             'Before using it in a business, we would check: ______________________']:
 doc.add_paragraph(text)
doc.add_paragraph('Remember: cycle service is a probability of no stockout; fill rate is the share of units served. A stock target is not the quantity to order.').runs[0].bold=True
doc.save(ROOT/'docs/T2_Decision_Sheet.docx')
